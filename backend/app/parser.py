from pathlib import Path
from io import BytesIO
import fitz
from docx import Document
from fastapi import HTTPException, UploadFile

async def parse(upload: UploadFile):
    name = Path(upload.filename or "resume.txt").name
    ext = Path(name).suffix.lower()
    if ext not in {".pdf", ".docx", ".txt"}:
        raise HTTPException(415, "Upload PDF, DOCX, or TXT.")
    data = await upload.read()
    if not data or len(data) > 8 * 1024 * 1024:
        raise HTTPException(400, "File is empty or exceeds 8 MB.")
    try:
        if ext == ".pdf":
            with fitz.open(stream=data, filetype="pdf") as d:
                text = "\n".join(p.get_text() for p in d)
        elif ext == ".docx":
            d = Document(BytesIO(data))
            text = "\n".join(p.text for p in d.paragraphs)
            for table in d.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows)
        else:
            text = data.decode("utf-8-sig")
    except Exception:
        raise HTTPException(400, "Could not read file. Try a text-based PDF or DOCX.")
    text = "\n".join(x.strip() for x in text.splitlines() if x.strip())
    if len(text) < 40:
        raise HTTPException(400, "Not enough readable text. Scanned PDFs need OCR.")
    return name, text
